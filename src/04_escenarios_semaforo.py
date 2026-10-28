"""
04_escenarios_semaforo.py — Proyección 2027/2030, escenarios, semáforo y base final.

Escenario 1 — SIN INTERVENCIÓN (inercia)
    Proyección del modelo (modelo.py) para la elección intermedia 2027 y la presidencial 2030.

Escenario 2 — CON INTERVENCIÓN HIPOTÉTICA (estrategia focalizada DECEyEC / ENCÍVICA)
    En las celdas objetivo (personas de 18-29 años en cualquier sección, o cualquier edad en
    secciones RURALES) la abstención proyectada baja un 15% relativo:  Â' = Â × (1 − 0.15).
    El resto de celdas no cambia. Se re-agrega a entidad con la lista nominal.
    · 2A: intervención en las 32 entidades.
    · 2B: intervención SOLO en entidades en rojo (versión con presupuesto restringido).
    · Sensibilidad: 5%, 10%, 15% y 20%.

Semáforo de riesgo (2027)
    Cuartiles de la abstención proyectada sin intervención. Los umbrales se fijan con el
    Escenario 1 y se reutilizan en el 2 para contar cuántas entidades cambian de color.

Capa operativa distrital (300 DEF, cartografía 2023)
    Abstención 2024 observada del distrito + cambio proyectado de su entidad (2024→2027).

Salidas
    data/processed/Base_Datos_Reto2_CCPC_2009_2024.csv / .xlsx   (32 filas, base maestra entregable)
    outputs/Modelo_Riesgo_Escenarios_Reto2.csv                  (resultados del modelo)
    outputs/proyeccion_celdas_2027.csv                          (entidad × sección × sexo × edad)
    outputs/semaforo_distritos_2027.csv                         (300 distritos)
    outputs/sensibilidad_intervencion.csv
"""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from config import (ENTIDADES, GOBERNATURA, OUT, PROC, RAW, REDUCCION_INTERVENCION,  # noqa: E402
                    SEGMENTO_OBJETIVO, SENSIBILIDAD, TIPOS_SECCION)
from modelo import agregar, matrices, proyectar, semaforo  # noqa: E402

ROJO, AMAR, VERDE = "Riesgo Alto (Rojo)", "Riesgo Medio (Amarillo)", "Riesgo Bajo (Verde)"


def lista_nominal_proyectada(L):
    """LN 2027 = LN 2024 × crecimiento trienal observado 2021→2024 de la entidad (aplicado a sus celdas)."""
    ent = L.index.get_level_values("EDOCVE")
    g = (L[2024].groupby(level="EDOCVE").sum() / L[2021].groupby(level="EDOCVE").sum())
    return L[2024] * g.reindex(ent).to_numpy()


def es_objetivo(idx):
    edad = idx.get_level_values("GRUPO_EDAD").isin(SEGMENTO_OBJETIVO["edades"])
    rural = idx.get_level_values("TIPOSEC").isin(SEGMENTO_OBJETIVO["tipos_seccion"])
    return edad | rural


def intervenir(p_sin, idx, reduccion, entidades=None):
    """Aplica la reducción relativa de abstención en celdas objetivo (y opcionalmente solo en ciertas entidades)."""
    m = es_objetivo(idx)
    if entidades is not None:
        m = m & idx.get_level_values("EDOCVE").isin(entidades)
    a = 100 - p_sin
    a = np.where(m, a * (1 - reduccion), a)
    return pd.Series(100 - a, index=idx)


def main():
    celdas = pd.read_csv(RAW / "ccpc_agregado_entidad.csv", encoding="utf-8-sig")
    P, L = matrices(celdas)

    # ------------------------------------------------------------------ Escenario 1 (inercia)
    P[2027], coefs27 = proyectar(P, L, 2027)
    L[2027] = lista_nominal_proyectada(L)
    P[2030], coefs30 = proyectar(P, L, 2030)          # usa la proyección 2027 como "memoria reciente"
    L[2030] = L[2027] * (L[2027].groupby(level="EDOCVE").transform("sum") / L[2024].groupby(level="EDOCVE").transform("sum"))

    a27 = 100 - agregar(P[2027], L[2027])
    a30 = 100 - agregar(P[2030], L[2030])
    sem, q_rojo, q_verde = semaforo(a27)

    # ------------------------------------------------------------------ Escenario 2 (intervención)
    p27_con = intervenir(P[2027], P.index, REDUCCION_INTERVENCION)
    a27_con = 100 - agregar(p27_con, L[2027])
    sem_con, _, _ = semaforo(a27_con, q_rojo, q_verde)
    rojos = sem[sem == ROJO].index.tolist()
    p27_rojos = intervenir(P[2027], P.index, REDUCCION_INTERVENCION, entidades=rojos)
    a27_rojos = 100 - agregar(p27_rojos, L[2027])
    sem_rojos, _, _ = semaforo(a27_rojos, q_rojo, q_verde)

    # Abstencionistas esperados (personas) y recuperados
    ln27 = L[2027].groupby(level="EDOCVE").sum()
    abst_pers_sin = (a27 / 100 * ln27).round()
    abst_pers_con = (a27_con / 100 * ln27).round()
    obj = es_objetivo(P.index)
    abst_obj = ((100 - P[2027]) / 100 * L[2027])[obj].groupby(level="EDOCVE").sum().round()

    # Error esperado: MAE de la validación retrospectiva del modelo final
    v = pd.read_csv(OUT / "validacion_backtest.csv")
    mae = float(v[v.Modelo.str.startswith("M5")].MAE_pp.mean())

    r = pd.DataFrame(index=a27.index)
    r["ENTIDAD"] = [ENTIDADES[i][0] for i in r.index]
    r["GUBERNATURA_2027"] = [int(i in GOBERNATURA[2027]) for i in r.index]
    r["LN_PROYECTADA_2027"] = ln27.round()
    r["Escenario_Sin_Intervencion_2027"] = a27.round(2)
    r["Banda_inferior_2027"] = (a27 - mae).round(2)
    r["Banda_superior_2027"] = (a27 + mae).round(2)
    r["Escenario_Sin_Intervencion_2030"] = a30.round(2)
    r["Semaforo_Riesgo"] = sem
    r["Escenario_Con_Intervencion"] = a27_con.round(2)
    r["Semaforo_Con_Intervencion"] = sem_con
    r["Cambia_color"] = (sem != sem_con).astype(int)
    r["Reduccion_abstencion_pp"] = (a27 - a27_con).round(2)
    r["Abstencionistas_esperados_sin"] = abst_pers_sin.astype(int)
    r["Abstencionistas_esperados_con"] = abst_pers_con.astype(int)
    r["Personas_recuperadas"] = (abst_pers_sin - abst_pers_con).astype(int)
    r["Abstencionistas_segmento_objetivo"] = abst_obj.astype(int)
    r["Escenario_2B_solo_rojos"] = a27_rojos.round(2)
    r["Semaforo_2B_solo_rojos"] = sem_rojos
    # Misma fórmula que 2A (diferencia de personas ya redondeadas por entidad) para que las cifras cuadren al dígito
    abst_pers_2b = (a27_rojos / 100 * ln27).round()
    r["Personas_recuperadas_2B"] = (abst_pers_sin - abst_pers_2b).astype(int)
    r["Ranking_riesgo"] = a27.rank(ascending=False, method="min").astype(int)

    # Partición de la proyección por segmento (para recomendaciones focalizadas)
    pc = pd.DataFrame({"P": P[2027], "L": L[2027]})
    for t, nom in TIPOS_SECCION.items():
        s = pc.xs(t, level="TIPOSEC")
        r[f"ABST_2027_{nom.upper()}"] = (100 - agregar(s.P, s.L)).round(2)
    juv = pc[pc.index.get_level_values("GRUPO_EDAD").isin(SEGMENTO_OBJETIVO["edades"])]
    r["ABST_2027_JUVENTUD_18_29"] = (100 - agregar(juv.P, juv.L)).round(2)
    for sx, nom in ((0, "HOMBRES"), (1, "MUJERES")):
        s = pc.xs(sx, level="SEXO")
        r[f"ABST_2027_{nom}"] = (100 - agregar(s.P, s.L)).round(2)
    r.index.name = "EDOCVE"
    r.to_csv(OUT / "Modelo_Riesgo_Escenarios_Reto2.csv", encoding="utf-8-sig")

    # ------------------------------------------------------------------ Celdas (insumo de tableros)
    cel = pd.DataFrame({"PART_2021": P[2021], "PART_2024": P[2024], "LN_2027": L[2027].round(),
                        "PART_2027_SIN": P[2027].round(2), "PART_2027_CON": p27_con.round(2),
                        "SEGMENTO_OBJETIVO": obj.astype(int)}).reset_index()
    cel["ENTIDAD"] = cel.EDOCVE.map(lambda i: ENTIDADES[i][0])
    cel["SEXO"] = cel.SEXO.map({0: "Hombres", 1: "Mujeres"})
    cel["TIPOSEC"] = cel.TIPOSEC.map(TIPOS_SECCION)
    cel.to_csv(OUT / "proyeccion_celdas_2027.csv", index=False, encoding="utf-8-sig")

    # ------------------------------------------------------------------ Sensibilidad
    sens = []
    for red in SENSIBILIDAD:
        a = 100 - agregar(intervenir(P[2027], P.index, red), L[2027])
        s, _, _ = semaforo(a, q_rojo, q_verde)
        nat = float((a * ln27).sum() / ln27.sum())
        sens.append({"Reduccion_relativa_segmento": f"{int(red * 100)}%",
                     "Abstencion_nacional_2027": round(nat, 2),
                     "Entidades_rojo": int((s == ROJO).sum()), "Entidades_amarillo": int((s == AMAR).sum()),
                     "Entidades_verde": int((s == VERDE).sum()), "Entidades_que_cambian_color": int((s != sem).sum()),
                     "Personas_recuperadas": int((abst_pers_sin - (a / 100 * ln27).round()).sum())})
    nat_sin = float((a27 * ln27).sum() / ln27.sum())
    sens.insert(0, {"Reduccion_relativa_segmento": "0% (sin intervención)", "Abstencion_nacional_2027": round(nat_sin, 2),
                    "Entidades_rojo": int((sem == ROJO).sum()), "Entidades_amarillo": int((sem == AMAR).sum()),
                    "Entidades_verde": int((sem == VERDE).sum()), "Entidades_que_cambian_color": 0, "Personas_recuperadas": 0})
    pd.DataFrame(sens).to_csv(OUT / "sensibilidad_intervencion.csv", index=False, encoding="utf-8-sig")

    # ------------------------------------------------------------------ Capa distrital 2024→2027
    d = pd.read_csv(RAW / "ccpc_agregado_distrito.csv")
    d = d[(d.AELEC == 2024) & d.DEF.notna()].copy()
    d["DEF"] = d.DEF.astype(int)
    d = d.groupby(["EDOCVE", "DEF"])[["LN", "SV", "NV"]].sum().reset_index()
    d["ABST_2024"] = 100 - 100 * d.SV / (d.SV + d.NV)
    a24 = 100 - agregar(P[2024], L[2024])
    d["CAMBIO_ENTIDAD_24_27"] = d.EDOCVE.map(a27 - a24)
    d["ABST_PROY_2027"] = (d.ABST_2024 + d.CAMBIO_ENTIDAD_24_27).clip(2, 95)
    d["SEMAFORO_DISTRITO"], _, _ = semaforo(d.ABST_PROY_2027)
    d["SEMAFORO_ENTIDAD"] = d.EDOCVE.map(sem)
    d["ABSTENCIONISTAS_ESPERADOS"] = (d.ABST_PROY_2027 / 100 * d.LN).round().astype(int)
    d.insert(1, "ENTIDAD", d.EDOCVE.map(lambda i: ENTIDADES[i][0]))
    d = d.round(2).sort_values("ABST_PROY_2027", ascending=False)
    d.to_csv(OUT / "semaforo_distritos_2027.csv", index=False, encoding="utf-8-sig")

    # ------------------------------------------------------------------ Base maestra entregable
    h = pd.read_csv(PROC / "historico_entidad.csv", index_col="EDOCVE", encoding="utf-8-sig")
    base = h.join(r.drop(columns=["ENTIDAD"])).copy()
    assert base.shape[0] == 32 and not base.isna().any().any(), "La base final debe tener 32 filas y 0 nulos"
    base.to_csv(PROC / "Base_Datos_Reto2_CCPC_2009_2024.csv", encoding="utf-8-sig")

    dicc = diccionario(base, set(r.columns) - {"ENTIDAD"})
    with pd.ExcelWriter(PROC / "Base_Datos_Reto2_CCPC_2009_2024.xlsx", engine="openpyxl") as xw:
        base.reset_index().to_excel(xw, sheet_name="Base_maestra_32_entidades", index=False)
        r.reset_index().to_excel(xw, sheet_name="Escenarios_y_semaforo", index=False)
        d.to_excel(xw, sheet_name="Semaforo_300_distritos", index=False)
        pd.DataFrame(sens).to_excel(xw, sheet_name="Sensibilidad", index=False)
        pd.read_csv(OUT / "validacion_backtest.csv").to_excel(xw, sheet_name="Validacion_modelo", index=False)
        pd.read_csv(OUT / "coeficientes_modelo.csv").to_excel(xw, sheet_name="Coeficientes", index=False)
        dicc.to_excel(xw, sheet_name="Diccionario", index=False)
    estilo_excel(PROC / "Base_Datos_Reto2_CCPC_2009_2024.xlsx")

    resumen = {"umbral_rojo": round(q_rojo, 2), "umbral_verde": round(q_verde, 2), "mae_validacion": round(mae, 2),
               "abstencion_nacional_2027_sin": round(nat_sin, 2),
               "abstencion_nacional_2027_con": round(float((a27_con * ln27).sum() / ln27.sum()), 2),
               # 2024 observado con la definición oficial de la CCPC: 100 − SV/(SV+NV) sobre el total nacional
               "abstencion_nacional_2024_obs": round(float(100 - 100 * celdas[celdas.AELEC == 2024].SV.sum()
                                                          / (celdas[celdas.AELEC == 2024].SV.sum() + celdas[celdas.AELEC == 2024].NV.sum())), 2),
               "coeficientes_2027": coefs27, "cambian_color": int(r.Cambia_color.sum()),
               "abstencion_nacional_2027_2B": round(float((a27_rojos * ln27).sum() / ln27.sum()), 2),
               "personas_recuperadas_2A": int(r.Personas_recuperadas.sum()), "personas_recuperadas_2B": int(r.Personas_recuperadas_2B.sum())}
    json.dump(resumen, open(OUT / "resumen_modelo.json", "w"), ensure_ascii=False, indent=2, default=str)
    print(json.dumps(resumen, ensure_ascii=False, indent=1, default=str))
    print(r[["ENTIDAD", "GUBERNATURA_2027", "Escenario_Sin_Intervencion_2027", "Semaforo_Riesgo",
             "Escenario_Con_Intervencion", "Semaforo_Con_Intervencion", "Personas_recuperadas"]]
          .sort_values("Escenario_Sin_Intervencion_2027", ascending=False).to_string())
    print(pd.DataFrame(sens).to_string(index=False))


def diccionario(base, cols_modelo):
    # ORDEN IMPORTA: next() devuelve la primera llave cuyo prefijo coincide, así que las llaves
    # específicas ("LN_PROYECTADA", "ABST_2027_") van antes que las genéricas ("LN_", "ABST_").
    desc = {
        "Personas_recuperadas_2B": "Abstencionistas evitados si se interviene solo en las entidades en rojo (Escenario 2B)",
        "Personas_recuperadas": "Abstencionistas evitados por la intervención nacional (Escenario 2A)",
        "LN_PROYECTADA": "Lista nominal proyectada a 2027 mediante tasa de crecimiento histórica (2021→2024)",
        "ABST_2027_": "Abstención proyectada 2027 en el segmento",
        "ENTIDAD": "Nombre de la entidad federativa", "ABREV": "Abreviatura",
        "LN_": "Lista nominal (CCPC) del año", "VOTANTES_": "Registros con marca de voto (SV)",
        "ABSTENCIONES_": "Registros sin marca de voto (NV)", "PART_": "Participación % = SV/(SV+NV)",
        "ABST_": "Abstención % = 100 − participación", "GOB_": "1 = gubernatura concurrente ese año",
        "CAIDA_INTERMEDIA": "Puntos de participación perdidos al pasar de presidencial a intermedia",
        "TENDENCIA_": "Cambio promedio anual (pp/año) entre elecciones del mismo tipo",
        "BRECHA_JUVENIL": "Participación 18-29 menos participación total (pp)",
        "BRECHA_GENERO": "Participación mujeres menos hombres (pp)",
        "BRECHA_RURAL": "Participación rural menos urbana (pp)",
        "PCT_LN_": "% de la lista nominal en el segmento", "PCT_HABLANTE": "Censo 2020 INEGI: % de 3+ años hablante de lengua indígena",
        "PCT_POB_HOGARES": "Censo 2020 INEGI: % de población en hogares indígenas",
        "PCT_POB_AFRO": "Censo 2020 INEGI: % de población afromexicana o afrodescendiente",
        "GUBERNATURA_2027": "1 = elige gubernatura en 2027",
        "Escenario_Sin_Intervencion_2027": "ESCENARIO 1: abstención % proyectada 2027 (inercia)",
        "Banda_": "Proyección ± error medio absoluto de la validación retrospectiva",
        "Escenario_Sin_Intervencion_2030": "Abstención % proyectada para la presidencial 2030 (inercia)",
        "Semaforo_Riesgo": "Semáforo 2027 sin intervención (cuartiles)",
        "Escenario_Con_Intervencion": "ESCENARIO 2: abstención % 2027 con intervención −15% en jóvenes y rural",
        "Semaforo_Con": "Semáforo con intervención (mismos umbrales)", "Cambia_color": "1 = cambia de color con la intervención",
        "Reduccion_abstencion_pp": "Reducción de abstención (pp) por la intervención",
        "Abstencionistas_esperados": "Personas que no votarían (proyección)",
        "Abstencionistas_segmento": "Abstencionistas esperados en jóvenes 18-29 o secciones rurales",
        "Escenario_2B": "Abstención 2027 si la intervención se aplica solo en entidades en rojo",
        "Semaforo_2B": "Semáforo si se interviene solo en rojos",
        "Ranking_riesgo": "1 = mayor abstención proyectada",
    }
    filas = []
    for col in base.reset_index().columns:
        d = next((v for k, v in desc.items() if col.startswith(k)), "Clave INE de entidad" if col == "EDOCVE" else "")
        if col.startswith("PART_") and col[5:].split("_")[0].isalpha():
            d = "Participación % del segmento en el año indicado"
        fuente = "Censo 2020 (INEGI)" if "2020" in col else ("Modelo del equipo" if col in cols_modelo else "CCPC 2009-2024 (INE)")
        filas.append({"Variable": col, "Descripción": d, "Fuente": fuente})
    return pd.DataFrame(filas)


def estilo_excel(ruta):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    wb = load_workbook(ruta)
    colores = {"Rojo": "F4B6B6", "Amarillo": "FFE9A8", "Verde": "BFE3C6"}
    for ws in wb.worksheets:
        ws.freeze_panes = "B2" if ws.title.startswith("Base") else "A2"
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="23395B")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(42, max(10, len(str(col[0].value)) + 2))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str):
                    for k, v in colores.items():
                        if f"({k})" in c.value:
                            c.fill = PatternFill("solid", fgColor=v)
    wb.save(ruta)


if __name__ == "__main__":
    main()
